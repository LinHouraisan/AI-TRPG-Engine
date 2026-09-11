from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "答辩材料"
OUT.mkdir(exist_ok=True)

BLUE = "2E74B5"
LIGHT = "E8EEF5"
GRAY = "F2F4F7"
INK = "17324D"
RED = "C00000"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_fixed_table(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths)-1)])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        hdr.cells[i].text = text
        shade(hdr.cells[i], LIGHT)
        for run in hdr.cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(INK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    set_fixed_table(table, widths or [9360 // len(headers)] * len(headers))
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_placeholder(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"【待替换：{text}】")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)
    return p


def configure(doc, short_title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in [("Title", 25, 0, 8), ("Heading 1", 16, 14, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11.5, 8, 4)]:
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(BLUE if name != "Title" else INK)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for list_name in ["List Bullet", "List Number"]:
        s = styles[list_name]
        s.font.name = "Microsoft YaHei"
        s.font.size = Pt(10.5)
        s.paragraph_format.space_after = Pt(3)
    header = section.header.paragraphs[0]
    header.text = f"AI TRPG Engine  |  {short_title}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(110, 120, 130)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("答辩材料  ·  2026年8月")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(110, 120, 130)


def cover(doc, title, subtitle):
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    kicker = doc.add_paragraph("AI TRPG ENGINE")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.runs[0].font.size = Pt(12)
    kicker.runs[0].font.bold = True
    kicker.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph(title, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(subtitle)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.color.rgb = RGBColor(70, 80, 90)
    doc.add_paragraph().paragraph_format.space_after = Pt(100)
    for line in ["组号：【待替换：组号】", "组长：【待替换：组长姓名】", "小组成员：【待替换：全部成员姓名】", "提交日期：2026年8月25日"]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "待替换" in line:
            p.runs[0].font.color.rgb = RGBColor.from_string(RED)
    doc.add_page_break()


def create_intro():
    doc = Document()
    configure(doc, "项目介绍")
    cover(doc, "项目介绍文档", "本地优先的 AI 单人跑团桌面引擎")
    doc.add_heading("1. 项目概况", level=1)
    add_table(doc, ["项目", "内容"], [
        ["项目名称", "AI TRPG Engine"],
        ["产品形态", "Windows 桌面应用（Electron）"],
        ["核心技术", "TypeScript、React、Electron、SQLite、可替换 AI 模型"],
        ["目标用户", "希望随时进行单人桌上角色扮演游戏的玩家"],
        ["当前阶段", "支持 AI 主持的单人跑团与完整试玩流程"],
    ], [1800, 7560])
    doc.add_heading("一句话价值", level=2)
    p = doc.add_paragraph()
    r = p.add_run("让玩家用自然语言自由行动，让 AI 负责主持和叙事，让程序负责规则、状态与事实一致性。")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)

    doc.add_heading("2. 业务场景", level=1)
    doc.add_paragraph("传统跑团依赖主持人、玩家和连续时间安排，单人玩家很难随时开始并长期持续。普通 AI 对话虽然能够生成叙事，却常把关键人物、道具、线索和既有结果遗忘或改写，难以形成可验证、可恢复的长期战役。")
    add_bullets(doc, [
        "玩家希望无需组织多人即可开始一场完整跑团。",
        "玩家希望使用自然语言表达行动，而不是只能点击固定选项。",
        "长期游戏需要保存世界状态、检定结果、剧情进度和关键事实。",
        "不同用户可能使用本地模型或自己的云端模型密钥。",
    ])
    add_placeholder(doc, "业务场景示意图；建议表现‘玩家自然语言输入—AI 主持—系统维护状态’")

    doc.add_heading("3. 核心功能", level=1)
    add_table(doc, ["功能模块", "用户能力", "产品价值"], [
        ["调查员创建", "分配属性并生成正式人物卡", "建立可持续使用的玩家角色"],
        ["剧本与战役", "选择内容包、创建并继续战役", "支持不同题材和重复游玩"],
        ["自然语言行动", "直接描述调查、交谈、移动或尝试", "减少固定菜单对自由度的限制"],
        ["规则与检定", "系统识别意图并执行百分制检定", "结果可验证、可追溯"],
        ["AI 主持叙事", "根据已提交结果生成场景反馈", "保持叙事体验与程序事实一致"],
        ["时间线与检查点", "查看历史、创建检查点、从副本恢复", "降低长期游玩丢档风险"],
        ["模型设置", "接入本地或云端模型并配置任务路由", "避免绑定单一模型供应商"],
    ], [1800, 3600, 3960])
    add_placeholder(doc, "系统主界面截图；建议包含叙事区、行动输入框、人物卡和时间线")

    doc.add_heading("4. 产品与案例特点", level=1)
    doc.add_heading("4.1 AI 与权威事实分离", level=2)
    doc.add_paragraph("系统采用“AI 提出候选，确定性系统提交事实”的原则。AI 可以理解玩家意图、提出变化并生成叙事，但不能直接修改战役数据库。规则裁定、随机结果、版本检查和状态提交由程序完成。")
    doc.add_heading("4.2 本地优先与可替换模型", level=2)
    doc.add_paragraph("战役数据保存在本机 SQLite 数据库中。模型可以接入本地服务，也可以使用用户自己的云端密钥；云端密钥通过 Electron 安全存储机制保存，不向渲染层暴露明文。")
    doc.add_heading("4.3 完整试玩切片", level=2)
    doc.add_paragraph("当前案例包含可运行的调查剧本、调查员创建、场景探索、人物交互、检定、线索与剧情推进等主流程，能够用于答辩现场展示产品的完整闭环。")
    add_placeholder(doc, "案例截图组：调查员创建、检定结果、剧情时间线，各放一张")

    doc.add_heading("5. 技术实现概览", level=1)
    add_table(doc, ["层级", "主要模块", "职责"], [
        ["交互层", "React Renderer", "展示战役、人物卡、叙事、检定和设置"],
        ["桌面桥接层", "Electron Main / Preload / IPC", "窗口生命周期、安全边界和接口调用"],
        ["运行时层", "Turn Router / Runtime / Keeper", "组织回合、选择处理路径、协调叙事"],
        ["确定性内核", "Rules / RNG / Validator / Events", "规则裁定、随机结果、校验与原子提交"],
        ["数据层", "SQLite / Event Log / Checkpoint", "保存权威状态、历史事件和恢复点"],
        ["模型接入层", "Provider / Profile / Task Route", "统一连接本地和云端模型"],
    ], [1500, 3000, 4860])
    add_placeholder(doc, "四层产品架构图；PPT 第3页架构图可复用")

    doc.add_heading("6. 当前成果与发展方向", level=1)
    add_bullets(doc, [
        "已经形成单一 Electron 桌面工作区，具备开发、测试、构建和 Windows 打包脚本。",
        "已接通战役、设置、回合、操作状态、时间线、检查点和备份等主要接口。",
        "已完成多个内容包和试玩案例，支持从人物创建进入正式游戏。",
        "后续将继续增强长期记忆、剧情调度、内容创作工具和多人协作能力。",
    ])
    doc.add_heading("7. 团队信息", level=1)
    add_table(doc, ["成员", "负责内容", "主要成果"], [
        ["【待替换：成员1】", "【待替换：职责】", "【待替换：成果】"],
        ["【待替换：成员2】", "【待替换：职责】", "【待替换：成果】"],
        ["【待替换：成员3】", "【待替换：职责】", "【待替换：成果】"],
    ], [2100, 3000, 4260])
    path = OUT / "AI TRPG Engine-项目介绍.docx"
    doc.save(path)
    return path


def create_spec():
    doc = Document()
    configure(doc, "功能规格说明书")
    cover(doc, "功能规格说明书（Spec）", "架构、数据库、API 与审查报告")
    doc.add_heading("文档控制", level=1)
    add_table(doc, ["字段", "内容"], [["版本", "V1.0（答辩版）"], ["状态", "评审稿"], ["适用产品", "AI TRPG Engine Desktop 0.1.0"], ["编制人", "【待替换：姓名】"]], [2200, 7160])
    doc.add_heading("目录", level=1)
    add_numbered(doc, ["范围与目标", "功能规格", "架构设计", "数据库设计", "API 接口设计", "安全与可靠性", "测试与验收", "审查报告"])

    doc.add_heading("1. 范围与目标", level=1)
    doc.add_heading("1.1 产品目标", level=2)
    doc.add_paragraph("构建一套本地优先、由 AI 主持的单人 TRPG 运行环境。玩家通过自然语言与世界互动；AI 负责理解、主持和叙事；程序负责规则、状态、因果关系和可恢复历史。")
    doc.add_heading("1.2 目标用户", level=2)
    add_bullets(doc, ["缺少固定主持人或队友的单人玩家。", "希望随时开始或继续战役的桌游爱好者。", "希望自行选择本地或云端模型的用户。"])
    doc.add_heading("1.3 非目标", level=2)
    add_bullets(doc, ["当前版本不实现多人实时联机。", "当前版本不提供真人 GM 工作台。", "AI 叙事文本不作为权威世界状态。", "不以大规模自主 NPC 社会模拟为核心目标。"])

    doc.add_heading("2. 功能规格", level=1)
    functions = [
        ("F-01 战役管理", "创建、列出、打开、关闭、移入回收站和恢复战役；显示健康状态、分支和状态版本。", "战役能够创建并重新打开，删除采用可恢复方式。"),
        ("F-02 调查员创建", "分配人物属性、确认人物卡，并在检查点后重新创建调查员。", "确认后生成稳定实体 ID、状态版本和检查点。"),
        ("F-03 行动提交", "玩家以自然语言提交行动，系统携带战役、分支、角色和期望状态版本。", "重复命令可识别；版本冲突不得覆盖新状态。"),
        ("F-04 回合裁定", "路由器区分询问、澄清和正式行动；规则引擎执行检定并生成事件。", "正式行动必须先裁定和提交，再生成正式叙事。"),
        ("F-05 AI 叙事", "根据已提交的公开结果生成主持人叙事，可使用模板或模型。", "叙事失败可重试，但不得重复掷骰或提交事件。"),
        ("F-06 时间线", "按游标分页读取战役历史，展示回合摘要与发生时间。", "历史顺序稳定，能够继续分页。"),
        ("F-07 检查点与恢复", "列出、创建检查点，并基于检查点创建恢复分支副本。", "恢复不得破坏原分支历史。"),
        ("F-08 模型设置", "管理 Provider、模型 Profile、任务路由和备用模型，并测试提供商。", "密钥不通过读取 API 返回明文。"),
        ("F-09 备份", "导出或导入带校验和、迁移信息和表数据的战役备份。", "导入前校验格式版本与校验和。"),
    ]
    add_table(doc, ["编号/模块", "规格描述", "验收标准"], functions, [1900, 4300, 3160])

    doc.add_heading("3. 架构设计", level=1)
    doc.add_heading("3.1 架构原则", level=2)
    add_bullets(doc, [
        "本地已提交状态和不可变事件是唯一权威事实源。",
        "AI 只能读取授权视图并输出结构化候选，不能直接写入权威数据库。",
        "规则裁定、随机结果、权限检查和状态提交由确定性程序负责。",
        "渲染进程不直接连接 SQLite，所有能力通过受控 IPC 暴露。",
        "模型接入与业务逻辑解耦，同一任务可配置主模型和备用模型。",
    ])
    doc.add_heading("3.2 物理架构", level=2)
    add_table(doc, ["组件", "代码位置", "职责"], [
        ["Renderer", "electron/src/renderer", "React 界面、交互状态和展示"],
        ["Preload", "electron/src/preload", "向 Renderer 暴露类型化桌面 API"],
        ["Main", "electron/src/main", "应用生命周期、IPC 注册、凭据与数据库组合"],
        ["Core Engine", "electron/src/core/engine", "路由、规则、检定、事件和回合运行"],
        ["Keeper", "electron/src/core/keeper", "AI 主持、上下文、流式输出和保护机制"],
        ["Persistence", "electron/sql + store", "SQLite 建表、迁移、状态与事件持久化"],
        ["Content Packs", "electron/content/packs", "剧本、地点、NPC、道具、事实和条件"],
    ], [1800, 3000, 4560])
    add_placeholder(doc, "架构图：Renderer → Preload/IPC → Main → Core Engine / Keeper → SQLite 与模型提供商")
    doc.add_heading("3.3 核心回合数据流", level=2)
    add_numbered(doc, [
        "玩家提交包含期望状态版本的自然语言行动。",
        "Turn Router 判断询问、澄清、自由行动或机械行动路径。",
        "运行时读取当前状态，执行规则、随机数和剧本条件。",
        "Validator 校验前后值、版本、权限和事件完整性。",
        "数据库原子提交状态与不可变事件。",
        "Keeper 根据已提交公开结果生成叙事并推送界面。",
    ])

    doc.add_heading("4. 数据库设计", level=1)
    doc.add_heading("4.1 设计目标", level=2)
    doc.add_paragraph("数据库同时保存可快速读取的当前状态和可追溯的不可变事件。删除、死亡、离场等生命周期变化通过状态或事件表达，不以物理删除历史记录代替。")
    add_table(doc, ["数据域", "代表数据", "约束"], [
        ["战役与分支", "campaign、branch、schema/version", "每个分支维护独立状态版本"],
        ["实体与状态", "调查员、NPC、场景、道具、事实、条件", "稳定 ID；引用目标必须存在且类型匹配"],
        ["事件日志", "回合事件、来源、序列、前后值", "追加写入；提交后不可静默修改"],
        ["检定记录", "技能、难度、骰值、结果", "随机种子与结果可追溯"],
        ["剧情进度", "节点、线索、调查和完成条件", "只有确定性运行时能够推进正式状态"],
        ["检查点", "分支、状态版本、事件序列、哈希、复盘", "恢复时创建副本，不覆盖原历史"],
        ["设置与模型", "Provider、Profile、任务路由", "凭据只保存引用 ID，不存明文"],
    ], [1800, 3400, 4160])
    doc.add_heading("4.2 关键一致性规则", level=2)
    add_bullets(doc, ["实体 ID 在战役内唯一且不复用。", "状态变更必须引用原因事件和前值版本。", "事件与状态在同一事务内原子提交。", "摘要、索引和活动上下文均可重建，不得成为新的事实源。", "叙事重试不得再次进行规则裁定或随机掷骰。"])
    add_placeholder(doc, "数据库 ER 图；建议突出 Campaign—Branch—Event—Entity State—Checkpoint 主关系")

    doc.add_heading("5. API 接口设计", level=1)
    doc.add_paragraph("桌面接口通过 Preload 暴露类型化 API，Renderer 不直接调用 Node.js、文件系统或 SQLite。API 当前主版本为 1。")
    api_rows = [
        ["app:getVersion / getState", "无", "版本或生命周期状态", "应用初始化与诊断"],
        ["campaign:create / list / open", "名称、分页或战役 ID", "CampaignSummary / View", "战役创建与访问"],
        ["campaign:confirmInvestigator", "战役、分支、属性分配", "人物卡、版本、检查点", "确认调查员"],
        ["turn:submitAction", "战役、分支、角色、版本、命令、文本", "operationId / turnId", "异步提交玩家行动"],
        ["operation:get / subscribe", "operationId、campaignId", "TurnView / 订阅 ID", "查询并接收回合进度"],
        ["timeline:page", "战役、分支、游标、数量", "分页时间线", "读取历史记录"],
        ["checkpoint:list / create", "战役、分支、标签和测试信息", "检查点或列表", "保存恢复点"],
        ["checkpoint:restoreCopy", "战役、检查点、标签", "新分支与版本", "从检查点恢复副本"],
        ["settings:listProviders / upsertProvider", "Provider 配置", "Provider 视图", "管理模型提供商"],
        ["settings:setSecret / hasSecret / deleteSecret", "凭据 ID 与密钥", "仅返回存在状态或 ID", "安全管理密钥"],
        ["backup:export / import", "战役 ID 或备份体", "备份或战役摘要", "战役迁移与恢复"],
    ]
    add_table(doc, ["接口", "关键输入", "输出", "用途"], api_rows, [2200, 3000, 2200, 1960])
    doc.add_heading("5.1 统一返回与错误处理", level=2)
    add_bullets(doc, ["所有异步接口返回统一 Result 结构。", "输入在 IPC 边界进行校验，非法输入不进入核心内核。", "长回合先返回 operationId，再通过查询或事件订阅获取结果。", "状态版本不一致时拒绝提交，避免旧页面覆盖新数据。", "模型失败与事实提交失败分开处理；已提交事实不因叙事失败回滚。"])

    doc.add_heading("6. 安全与可靠性", level=1)
    add_table(doc, ["风险", "控制措施"], [
        ["渲染进程越权", "仅通过 Preload 白名单 API 与类型化 IPC 通道访问能力"],
        ["密钥泄露", "使用 Electron safeStorage 加密保存；Renderer 无 getSecret 接口"],
        ["AI 篡改事实", "AI 仅生成候选或叙事；Core Engine 负责校验与提交"],
        ["重复操作", "commandId、operationId 和回合 ID 用于识别与追踪"],
        ["并发覆盖", "提交时检查 expectedStateVersion"],
        ["叙事失败", "可依据同一提交结果重试，不重复裁定和掷骰"],
        ["存档损坏", "检查点、状态哈希、备份校验和与恢复分支"],
    ], [2500, 6860])

    doc.add_heading("7. 测试与验收", level=1)
    add_table(doc, ["测试层级", "验证内容", "建议命令/证据"], [
        ["静态检查", "生产源码类型正确", "bun run typecheck"],
        ["单元测试", "输入校验、规则、检定、Keeper、存储", "bun run test"],
        ["持久化检查", "DDL、迁移、事件不可变性、凭据", "bun run persist:check"],
        ["架构检查", "模块依赖与边界符合约束", "bun run architecture:check"],
        ["核心流程", "人物创建、行动、检定、叙事和时间线", "bun run demo:e2e / 现场黄金路径"],
        ["构建验收", "Renderer 与 Main 可构建", "bun run build"],
        ["Windows 交付", "win-unpacked 可启动", "bun run package:win"],
    ], [1900, 4300, 3160])

    doc.add_heading("8. 审查报告", level=1)
    add_table(doc, ["审查项", "审查证据", "结论", "遗留事项"], [
        ["功能覆盖", "主要界面、回合、时间线、检查点和设置已有实现与测试", "基本通过", "内容创作与多人功能不在当前范围"],
        ["架构一致性", "Renderer 不直连 SQLite；Core 负责权威提交", "通过", "继续保持依赖边界"],
        ["数据库设计", "状态、事件、检查点和迁移脚本并存", "通过", "正式提交前补充最终 ER 图"],
        ["API 设计", "共享 api.ts 定义白名单通道与输入输出类型", "通过", "预留 content/model 列表接口仍待完整实现"],
        ["安全性", "凭据走 safeStorage，Renderer 无明文读取能力", "通过", "需在目标答辩机器验证安全存储可用"],
        ["可靠性", "版本校验、原子提交、检查点、备份与测试脚本", "基本通过", "现场模型响应仍受本地环境影响"],
        ["演示准备", "可使用固定试玩案例完成黄金路径", "有条件通过", "需准备预加载存档和备用录屏"],
    ], [1700, 3500, 1300, 2860])
    doc.add_heading("8.1 审查结论", level=2)
    doc.add_paragraph("AI TRPG Engine 已具备答辩演示所需的产品闭环和技术证据。架构、数据库与 API 的关键边界清晰，能够说明项目不是单纯的 AI 聊天界面，而是具备确定性规则、权威状态和可恢复历史的桌面运行环境。建议提交前完成占位信息替换、ER 图补充、目标机器演示验证和备用录屏。")
    doc.add_heading("附录 A：答辩前检查清单", level=1)
    add_bullets(doc, ["替换组号、组长、成员、分工等全部红色占位符。", "补入架构图、ER 图和系统截图。", "在答辩电脑上验证字体、PPT、程序和视频均可打开。", "使用固定存档演练三次，正式讲解控制在 6 分钟。", "准备 PPT PDF 备份和 90—120 秒系统录屏。"])
    path = OUT / "AI TRPG Engine-功能规格说明书.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(create_intro())
    print(create_spec())
