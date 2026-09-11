from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "答辩材料" / "AI TRPG Engine-操作手册.docx"
ARCH = ROOT / "答辩材料" / "images" / "system-architecture.png"
ER = ROOT / "答辩材料" / "images" / "database-er-diagram.png"

NAVY = "0B2545"
BLUE = "1769D2"
CYAN = "12B8B0"
LIGHT = "EAF2FA"
PALE = "F5F8FB"
GRAY = "667085"
LINE = "D5DEE8"
RED = "A43B3B"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for i, width in enumerate(widths_cm):
            row.cells[i].width = Cm(width)
            set_cell_margins(row.cells[i])
            set_cell_border(row.cells[i])
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(Cm(width).twips)))
        grid.append(col)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(int(Cm(w).twips) for w in widths_cm)))
    tbl_w.set(qn("w:type"), "dxa")


def set_run(run, size=10.5, bold=False, color=NAVY, font="Microsoft YaHei"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, bold_lead=None, style=None, after=5, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run(r2)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def add_note(doc, title, text, tone="info"):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [16.0])
    cell = table.cell(0, 0)
    shade(cell, "FFF4E5" if tone == "warn" else LIGHT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run(r, bold=True, color=RED if tone == "warn" else BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    set_run(p2.add_run(text), size=10, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_steps(doc, steps):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "space")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(tabs)
    ppr.append(ind)
    lvl.extend([start, num_fmt, lvl_text, suff, ppr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)

    for step in steps:
        p = doc.add_paragraph()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_el])
        p._p.get_or_add_pPr().append(num_pr)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.2
        set_run(p.add_run(step), size=10.5)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        set_run(p.add_run(item), size=10.5)


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    set_run(p.add_run(text), size=17, bold=True, color=BLUE)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    set_run(p.add_run(text), size=13, bold=True, color=NAVY)
    return p


def add_figure(doc, path, caption, width_cm=15.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(7)
    set_run(c.add_run(caption), size=9, color=GRAY)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 17, BLUE, 18, 10),
        ("Heading 2", 13, NAVY, 13, 7),
        ("Heading 3", 11.5, NAVY, 9, 5),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run(run, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_run(run2, size=9, color=GRAY)


def build():
    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.8)

    hp = sec.header.paragraphs[0]
    set_run(hp.add_run("AI TRPG Engine  |  用户操作手册"), size=9, color=GRAY)
    add_page_number(sec.footer.paragraphs[0])

    # Cover: editorial_cover with an A4 named override for Chinese defense materials.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("AI TRPG Engine"), size=30, bold=True, color=NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(16)
    set_run(p2.add_run("操 作 手 册"), size=22, bold=True, color=BLUE)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p3.add_run("《雾港末班车》单人 AI 跑团 Demo"), size=13, color=CYAN)
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_before = Pt(78)
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p4.add_run("适用平台：Windows x64  |  应用版本：0.1.0"), size=10, color=GRAY)
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p5.add_run("编制日期：2026 年 8 月"), size=10, color=GRAY)

    doc.add_page_break()
    add_h2(doc, "使用说明")
    add_text(doc, "本手册依据项目当前 Electron 代码、随包资料和 docs 设计文档编写。可操作步骤以当前 Demo 已实现界面为准；docs 中标记为 Draft、但尚未出现在当前界面的 V1 设计不作为可用功能介绍。")
    add_note(doc, "数据与模型说明", "战役数据默认保存在本机 SQLite。使用 DeepSeek 时，完成当前任务所需的受限上下文会发送到 DeepSeek；API Key 由系统安全存储加密保存，不写入战役数据库或战役备份。")
    add_h2(doc, "目录")
    toc = [
        "1. 系统简介与运行环境", "2. 安装与启动", "3. 首次使用与 DeepSeek 模型配置",
        "4. 新建战役与选择《雾港末班车》", "5. 创建调查员", "6. 主界面功能区域说明",
        "7. 输入行动、推进剧情与查看检定", "8. 人物卡、地图、时间线与上下文信息",
        "9. 创建检查点、复制恢复与分支管理", "10. 战役备份、导入和重新打开",
        "11. 常见问题、使用限制与安全说明", "12. 附录：推荐答辩演示流程",
    ]
    add_bullets(doc, toc)

    add_h1(doc, "1. 系统简介与运行环境")
    add_text(doc, "AI TRPG Engine 是一款本地优先的单人桌上角色扮演游戏引擎。玩家使用自然语言描述行动，AI 守秘人负责叙述和扮演 NPC；规则判定、随机结果、权威状态与存档由程序负责。")
    add_h2(doc, "1.1 当前 Demo 能力")
    add_bullets(doc, [
        "运行《雾港末班车》单人调查模组，支持九个场景、五名 NPC、十五件道具、十五条线索与四种结局路径。",
        "支持建议行动和自然语言自由行动；需要检定时由程序选择技能、难度并生成可追溯结果。",
        "支持调查员创建、人物卡、当前场景、背包、线索、剧情标记、时间线、检查点、分支和战役备份。",
        "正式云模型配置以 DeepSeek 为准；模型不可用时部分确定性模板流程仍可运行，但自由文本理解和正式 GM 叙述受限。",
    ])
    add_h2(doc, "1.2 运行环境")
    table = doc.add_table(rows=1, cols=2)
    set_table_widths(table, [4.3, 11.7])
    for i, value in enumerate(("项目", "要求或说明")):
        shade(table.rows[0].cells[i], NAVY)
        set_run(table.rows[0].cells[i].paragraphs[0].add_run(value), bold=True, color=WHITE)
    rows = [
        ("操作系统", "Windows x64"),
        ("启动方式", "安装版，或 win-unpacked 目录中的 AI TRPG Engine.exe"),
        ("网络", "本地存档无需联网；调用 DeepSeek 需要网络连接"),
        ("模型凭据", "用户自行准备 DeepSeek API Key；可能产生模型调用费用"),
        ("数据位置", "战役与设置由桌面主进程管理；战役数据使用 SQLite"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        shade(cells[0], PALE)
        set_run(cells[0].paragraphs[0].add_run(left), bold=True)
        set_run(cells[1].paragraphs[0].add_run(right))

    add_figure(doc, ARCH, "图 1  系统运行架构（Renderer 通过 Preload/IPC 访问主进程与核心引擎）")

    add_h1(doc, "2. 安装与启动")
    add_h2(doc, "2.1 使用打包版本")
    add_steps(doc, [
        "取得 Windows x64 安装包或解压后的 win-unpacked 文件夹。",
        "安装版按向导完成安装；解压版直接打开 AI TRPG Engine.exe。",
        "等待主窗口出现。首次打开时，应用会创建本地设置与战役存储。",
        "确认顶部显示 Electron、战役目录和存档状态；若按钮暂时不可用，请等待当前初始化完成。",
    ])
    add_h2(doc, "2.2 开发环境启动（仅项目维护者）")
    add_text(doc, "在仓库根目录安装依赖后，可使用 bun run desktop 启动桌面程序。Electron 工作区也提供 bun run dev、bun run typecheck、bun run test 和 bun run build 等维护命令。普通用户不需要执行这些命令。")
    add_note(doc, "注意", "不要把 node_modules、源码目录或开发命令当作最终用户安装步骤。答辩演示优先使用已打包的 Windows 程序。", "warn")

    add_h1(doc, "3. 首次使用与 DeepSeek 模型配置")
    add_h2(doc, "3.1 配置步骤")
    add_steps(doc, [
        "单击窗口顶部的“设置”。",
        "单击“使用 DeepSeek 云端”，系统会写入 DeepSeek 的预设地址和当前模型配置。",
        "在 API Key 输入框粘贴自己的 DeepSeek API Key，然后单击保存。保存成功后，输入框不会回显原始密钥。",
        "单击“测试连接”。等待系统检查模型是否存在、最小生成是否成功以及 JSON 输出是否通过。",
        "看到“连接成功”后关闭设置面板，再开始正式自由行动。",
    ])
    add_h2(doc, "3.2 配置状态说明")
    add_bullets(doc, [
        "“API Key 已由系统加密保存”：密钥已通过 Electron safeStorage 保护，并由凭据标识关联到模型设置。",
        "“模型不存在”或“生成失败”：检查模型名称、账户权限、网络状态和账户余额。",
        "“连接失败”：检查 API Key 是否有效，网络是否能访问 DeepSeek，以及服务是否暂时限流。",
        "模型用量区域显示调用次数、输入与输出 Token 及估算值；实际费用以 DeepSeek 平台账单为准。",
    ])
    add_note(doc, "隐私边界", "只有完成当前 GM 任务所需的受限上下文会发送给模型提供商。战役导出文件不包含模型设置和 API Key。")

    add_h1(doc, "4. 新建战役与选择《雾港末班车》")
    add_h2(doc, "4.1 新建战役")
    add_steps(doc, [
        "在顶部战役栏单击“新建”。",
        "系统创建名为“未命名战役”的本地战役，并自动切换到该战役。",
        "从下拉列表可切换已存在的战役；右侧会显示当前共有多少场战役。",
        "如需删除当前战役，单击“删除”并在确认框中确认。删除属于高风险操作，建议先导出备份。",
    ])
    add_h2(doc, "4.2 选择模组")
    add_steps(doc, [
        "单击左上角的模组标题区域，打开模组列表。",
        "选择《雾港末班车》。若当前已在使用该模组，条目会显示“正在用”。",
        "确认切换。页面会重新加载，并按目标资料包重新装配内容。",
    ])
    add_note(doc, "模组与战役的关系", "换模组等于离开当前时间线。当前战役会保留在数据库中；目标模组若已有存档则继续，没有则从新战役开始。不同资料包的事件不会混写到同一条时间线。")

    add_h1(doc, "5. 创建调查员")
    add_text(doc, "《雾港末班车》在正式开场前要求完成五步创建：前提、职业、技能、经历和确认。调查员一经确认即正式开局，普通分支中不可编辑。")
    add_h2(doc, "5.1 填写姓名与查看固定属性")
    add_steps(doc, [
        "阅读“前提”页的故事简介，单击“下一步”。",
        "在“职业”页输入调查员姓名。职业固定为“记者”。",
        "查看八项固定属性：STR 50、CON 50、SIZ 60、DEX 60、APP 50、INT 70、POW 65、EDU 70。",
    ])
    add_h2(doc, "5.2 分配技能点")
    add_text(doc, "需要恰好分完 280 点职业点和 140 点兴趣点。职业点只能投入职业技能；兴趣点可以投入全部列出的技能。单项最终技能不得超过 90。")
    table = doc.add_table(rows=1, cols=3)
    set_table_widths(table, [5.4, 4.0, 6.6])
    for i, value in enumerate(("技能", "基础值", "职业点可投入")):
        shade(table.rows[0].cells[i], NAVY)
        set_run(table.rows[0].cells[i].paragraphs[0].add_run(value), bold=True, color=WHITE)
    for row in [("侦查", "25", "是"), ("聆听", "20", "是"), ("图书馆使用", "20", "是"), ("话术", "5", "是"), ("心理学", "10", "是"), ("开锁", "1", "否，仅可投入兴趣点")]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if i == 0: shade(cells[i], PALE)
            set_run(cells[i].paragraphs[0].add_run(value), bold=(i == 0))
    h = add_h2(doc, "5.3 选择人生经历并确认")
    h.paragraph_format.page_break_before = True
    add_bullets(doc, [
        "档案通信员：偏向核对旧档案，并获得与沈鹭及档案相关的初始信息。",
        "旧线事故记者：与旧线事故报道和顾弦相关。",
        "站务账簿抄写员：与站务账簿、罗姨和名单核对相关。",
        "潮汐合影摄影师：与站员合影和神秘女孩相关。",
    ])
    add_steps(doc, [
        "选择一项人生经历。当前版本不支持自由输入背景。",
        "进入“确认”页，核对姓名、固定属性、生命、理智、技能和关系。",
        "确认职业点与兴趣点剩余均为 0，且页面没有校验错误。",
        "单击“确认调查员并开始”。系统写入不可变调查员档案并显示正式开场。",
    ])

    add_h1(doc, "6. 主界面功能区域说明")
    add_text(doc, "桌面宽屏采用三栏主持桌。叙事位于中心，角色与场景信息位于左侧，记录与战役状态位于右侧。窄窗口会切换为底部栏目导航。")
    table = doc.add_table(rows=1, cols=3)
    set_table_widths(table, [3.2, 4.2, 8.6])
    for i, value in enumerate(("区域", "主要组件", "用途")):
        shade(table.rows[0].cells[i], NAVY)
        set_run(table.rows[0].cells[i].paragraphs[0].add_run(value), bold=True, color=WHITE)
    for row in [
        ("顶部", "战役、模组、设置、恢复、导入导出", "管理当前战役、模型和存档工具。"),
        ("左栏", "调查员卡、当前场景", "查看属性、技能、生命、理智、当前位置、已访问地点与在场 NPC。"),
        ("中栏", "GM 叙事、建议行动、自由输入", "阅读剧情、选择建议或输入自己的行动。"),
        ("右栏", "背包、线索、剧情标记、时间线、事件记录", "查看已提交的权威状态与历史。"),
        ("移动/窄窗", "叙述、调查员、场景、记录", "通过底部四个栏目在同一信息结构间切换。"),
    ]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if i == 0: shade(cells[i], PALE)
            set_run(cells[i].paragraphs[0].add_run(value), bold=(i == 0))
    add_note(doc, "权威状态", "右侧显示的背包、线索、剧情标记、事件和版本来自已提交状态。GM 叙述不能直接覆盖这些事实。")

    add_h1(doc, "7. 输入行动、推进剧情与查看检定")
    add_h2(doc, "7.1 两种行动方式")
    add_steps(doc, [
        "阅读 GM 当前叙述和建议行动。",
        "快速操作：直接单击一个建议行动。建议行动携带结构化意图，通常处理更快、更稳定。",
        "自由操作：在输入框用自然语言描述想做的事，然后单击“行动”或按 Enter。Shift+Enter 可换行。",
        "等待“正在处理”状态结束。处理期间不要重复提交同一行动。",
        "阅读 GM 返回的新叙述、规则卡片和下一组建议行动。",
    ])
    add_text(doc, "输入框提示为“你打算做什么？直接说就行，建议行动只是省事，不是能做的全部。”玩家无需输入 /roll，也无需自己指定检定技能或难度。")
    add_h2(doc, "7.2 检定结果")
    add_bullets(doc, [
        "系统根据规则、当前场景和行动条件判断是否需要检定。",
        "公开检定会显示采用的技能值、普通/困难/极难阈值、骰点与结果。",
        "人物卡会高亮当前相关技能，时间线会标记该回合“掷了骰”。",
        "“换一种说法”只基于同一批已提交事件重新叙述，不重掷骰、不改变状态版本。",
    ])
    add_note(doc, "推荐演示输入", "在七号站台输入：“查看四周，重点注意是否有可以让列车员回答问题的办法”。验收资料预期人物卡高亮侦查 87，并显示可追溯检定结果。")

    add_h1(doc, "8. 人物卡、地图、时间线与上下文信息")
    add_h2(doc, "8.1 调查员卡与当前场景")
    add_bullets(doc, [
        "调查员卡显示已确认的档案、固定属性、生命、理智、技能和人生经历关系。",
        "当前场景显示团内时间、调查员所在地点、已访问地点、可见道具数量和在场 NPC。未访问地点不会泄露真实名称。",
        "背包列出调查员当前持有物品；已知线索只显示已经确认的事实。",
    ])
    add_h2(doc, "8.2 时间线、分支和事件记录")
    add_bullets(doc, [
        "时间线按回合聚合事件，每个条目显示状态版本、摘要、团内时间和掷骰标记。",
        "单击“回到这一版”不会删除之后的内容，而是从该版本创建新分支。",
        "出现多个分支后，可在分支列表中单击“切过去”查看或继续另一条时间线。",
        "事件记录只追加，显示玩家可见的权威事件及版本。",
    ])
    add_h2(doc, "8.3 上下文与调试信息")
    add_text(doc, "上下文使用面板用于说明当前角色、场景、线索等信息如何占用上下文预算。调试后台任务默认关闭；即使开启，也只显示 after-commit 的只读诊断，不应将幕后秘密当作玩家信息。")

    add_h1(doc, "9. 创建检查点、复制恢复与分支管理")
    add_h2(doc, "9.1 创建检查点")
    add_steps(doc, [
        "单击顶部“测试与恢复”。",
        "单击“创建检查点”。系统以当前分支和状态版本创建手动检查点。",
        "检查点条目显示标签、版本、前情摘要和状态哈希。",
    ])
    add_h2(doc, "9.2 复制并恢复")
    add_steps(doc, [
        "在检查点列表中找到目标检查点，单击“复制并恢复”。",
        "如果目标版本早于当前进度，阅读确认提示并选择是否继续。",
        "系统从检查点创建新的恢复分支并自动进入该分支；来源分支、事件和叙述保持不变。",
        "恢复视图会显示检查点前情和近期真实对话，检查点之后的回合不会出现在恢复分支中。",
    ])
    add_h2(doc, "9.3 重新创建调查员")
    add_text(doc, "只有标签为“正式开局前”的检查点提供“重新创建调查员”。使用后会建立一条未开局子分支，让用户重新完成调查员创建；来源分支仍绑定原调查员，不会被修改。")
    add_note(doc, "分支原则", "恢复和回到旧版本均采用“复制后继续”，而不是覆盖原时间线。这样可以尝试不同选择，同时保留来源分支。")

    add_h1(doc, "10. 战役备份、导入和重新打开")
    add_h2(doc, "10.1 导出战役")
    add_steps(doc, [
        "切换到需要备份的战役。",
        "单击顶部“导出”。",
        "选择保存位置并保管导出的 JSON 文件。导出内容是战役权威数据，不是单纯聊天记录。",
    ])
    add_h2(doc, "10.2 导入战役")
    add_steps(doc, [
        "单击顶部“导入”，选择本引擎导出的 JSON 文件。",
        "系统校验备份格式、数据完整性和资料包引用。",
        "校验通过后导入并加载战役；若资料包版本不匹配或数据被篡改，系统会拒绝导入。",
    ])
    add_h2(doc, "10.3 关闭后重新打开")
    add_text(doc, "正常关闭应用后再次启动，从顶部战役下拉列表选择原战役即可继续。当前战役、分支、调查员、事件和叙述会由本地数据库恢复。")
    add_note(doc, "备份不包含的内容", "战役备份不包含桌面设置、模型配置和 API Key。更换电脑后需重新配置 DeepSeek 凭据。", "warn")
    add_figure(doc, ER, "图 2  战役、分支、事件、实体状态与检查点的主要数据关系", width_cm=12.8)

    add_h1(doc, "11. 常见问题、使用限制与安全说明")
    qa = [
        ("模型连接失败怎么办？", "重新打开“设置”，检查 DeepSeek API Key、网络、模型名称、账户权限和余额，然后再次测试连接。"),
        ("为什么不能修改调查员？", "调查员确认后即绑定当前分支。若要重建，使用“正式开局前”检查点的“重新创建调查员”。"),
        ("为什么恢复后原进度还在？", "系统恢复时创建新分支，来源分支不会被覆盖。可在时间线分支列表中切换。"),
        ("为什么自由输入效果受限？", "正式自由文本路由与 GM 叙述依赖可用的 DeepSeek 配置；模型不可用时只能使用部分确定性流程或模板。"),
        ("导入为什么被拒绝？", "备份必须由本引擎导出，并通过结构、哈希和资料包引用校验。被篡改或资料包不匹配的文件会被拒绝。"),
        ("数据是否全部离线？", "战役状态和存档本地保存；调用 DeepSeek 时，当前任务所需的受限上下文会发送到云端。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_widths(table, [5.8, 10.2])
    for i, value in enumerate(("问题", "处理方式")):
        shade(table.rows[0].cells[i], NAVY)
        set_run(table.rows[0].cells[i].paragraphs[0].add_run(value), bold=True, color=WHITE)
    for question, answer in qa:
        cells = table.add_row().cells
        shade(cells[0], PALE)
        set_run(cells[0].paragraphs[0].add_run(question), bold=True)
        set_run(cells[1].paragraphs[0].add_run(answer), size=10)
    add_h2(doc, "11.1 当前版本限制")
    add_bullets(doc, [
        "当前是 Demo，不是完整 V1.0；没有通用作者工具、自动更新、多人模式或真人 GM 模式。",
        "当前调查员职业和属性固定，人生经历只有四种，不能自由输入背景。",
        "正式云模型只支持 DeepSeek；Ollama 兼容代码保留，但不是当前默认认证配置。",
        "Windows 打包可运行，但仍使用 Electron 默认图标，部分发布包装信息尚未完善。",
    ])

    add_h1(doc, "12. 附录：推荐答辩演示流程")
    add_text(doc, "以下流程覆盖模型配置、调查员创建、自然语言行动、检定、检查点、恢复和分支，是展示当前系统核心价值的最短闭环。")
    add_steps(doc, [
        "启动 Windows 程序，展示顶部战役栏、本地存档状态和三栏主持桌。",
        "打开“设置”，说明 DeepSeek API Key 由系统加密保存，并完成连接测试。",
        "新建战役并选择《雾港末班车》。",
        "创建调查员：分完 280 职业点、140 兴趣点，选择“潮汐合影摄影师”，确认开局。",
        "在七号站台输入推荐观察语句，展示自然语言路由、侦查检定、人物卡高亮和可追溯结果。",
        "推进至少三轮玩家/GM 对话，创建检查点，再推进一轮。",
        "使用“复制并恢复”，展示恢复前情、最近对话及新分支；切回来源分支说明原时间线未被覆盖。",
        "最后展示导出/导入入口、架构图和 ER 图，说明本地优先、事件溯源和模型隔离。",
    ])
    add_note(doc, "演示前检查", "使用全新战役；提前确认 DeepSeek 可用且账户余额充足；不要使用包含剧透或开发测试数据的旧分支；准备一份已导出的战役备份用于演示导入。")

    add_h2(doc, "文档依据")
    add_bullets(doc, [
        "项目 README 与 electron/README.md",
        "Electron Renderer：App、ModelSettings、CampaignDock、PackSelector、InvestigatorCreation、Composer、Timeline、CheckpointTests 等组件",
        "docs/00-product/player-experience.md、docs/01-architecture/game-loop.md、docs/02-data/save-branch.md",
        "docs/05-implementation-design/12-model-providers.md、15-player-ui.md、18-release-compatibility.md",
        "docs/demo/mist-harbor-test-cases.md 与 mist-harbor-known-issues.md",
    ])

    doc.core_properties.title = "AI TRPG Engine 操作手册"
    doc.core_properties.subject = "《雾港末班车》单人 AI 跑团 Demo 使用说明"
    doc.core_properties.author = "AI TRPG Engine 项目组"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
